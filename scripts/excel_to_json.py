"""
Конвертер Excel → bloggers.json

Использование:
    python scripts/excel_to_json.py

Требует: openpyxl (pip install openpyxl)
"""
import json
import re
from pathlib import Path
from openpyxl import load_workbook

EXCEL_PATH = Path("самая лучшая соцсеть!!!!.xlsx")
OUTPUT_PATH = Path("data/sources/bloggers.json")
SHEET_NAME = "Лист4"

UPDATED_AT = "2026-09-28"


def extract_id(url: str) -> str:
    match = re.search(r"max\.ru/(?:join/)?([^/?#]+)", url)
    return match.group(1) if match else url


def parse_ads(value) -> bool | None:
    if value is None:
        return None
    v = str(value).strip().lower()
    if v in ("да", "yes", "true", "1"):
        return True
    if v in ("нет", "no", "false", "0"):
        return False
    return None


def parse_topics(value) -> list[str]:
    """Парсит категории через запятую → список без дублей."""
    if not value:
        return []
    parts = str(value).split(",")
    seen = []
    for p in parts:
        p = p.strip().lower()
        if p and p not in seen:
            seen.append(p)
    return seen


def main():
    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    headers = [cell.value for cell in ws[1]]
    col_idx = {h: i for i, h in enumerate(headers) if h}

    bloggers = []
    seen_urls = set()

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue

        url = row[col_idx["ссылка"]]
        if not url:
            print(f"⚠️  Строка {row_num}: пустой URL, пропущено")
            continue
        if url in seen_urls:
            print(f"⚠️  Строка {row_num}: дубликат {url}, пропущено")
            continue
        seen_urls.add(url)

        name_raw = row[col_idx.get("название", -1)] if "название" in col_idx else None
        name = str(name_raw).strip() if name_raw else f"Канал {extract_id(url)}"

        blogger = {
            "id": extract_id(url),
            "name": name,
            "url": url,
            "topics": parse_topics(row[col_idx.get("категория", -1)] if "категория" in col_idx else None),
            "description": "",
            "subscribers": int(row[col_idx["подписчики"]] or 0),
            "posts_total": int(row[col_idx["публикаций"]] or 0),
            "posts_last_30d": None,
            "avg_views": int(row[col_idx["ср. просмотры"]] or 0),
            "avg_likes": int(row[col_idx["ср. реакции"]] or 0),
            "avg_comments": None,
            "has_ads": parse_ads(row[col_idx.get("есть реклама?", -1)] if "есть реклама?" in col_idx else None),
            "ads_count_30d": None,
            "updated_at": UPDATED_AT,
        }
        bloggers.append(blogger)

    bloggers.sort(key=lambda b: b["subscribers"], reverse=True)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(bloggers, f, ensure_ascii=False, indent=2)

    print(f"\n✅ Готово: {len(bloggers)} блогеров → {OUTPUT_PATH}")

    with_ads = sum(1 for b in bloggers if b["has_ads"] is True)
    without_ads = sum(1 for b in bloggers if b["has_ads"] is False)
    unknown_ads = sum(1 for b in bloggers if b["has_ads"] is None)
    print(f"   📢 С рекламой: {with_ads}")
    print(f"   🚫 Без рекламы: {without_ads}")
    print(f"   ❓ Неизвестно: {unknown_ads}")

    # Список всех тем — пригодится для FSM
    all_topics = sorted({t for b in bloggers for t in b["topics"]})
    print(f"\n📚 Уникальные темы ({len(all_topics)}):")
    for t in all_topics:
        print(f"   • {t}")


if __name__ == "__main__":
    main()